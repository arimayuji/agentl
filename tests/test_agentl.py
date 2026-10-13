import json
import urllib.request

import pytest

from agentl import SpecError, connector, graph, parse, register_provider
from agentl.llm import Result, schema_model
from agentl.runner import Runner
from agentl.spec import load
from agentl.store import Store


@pytest.fixture
def store(tmp_path):
    with Store(tmp_path / "db.sqlite") as s:
        yield s


def test_example_pipeline_runs(store, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path.parent)  # out/ goes to a temp dir
    import pathlib
    import shutil

    root = pathlib.Path(__file__).parents[1]
    shutil.copytree(root / "examples", tmp_path.parent / "examples", dirs_exist_ok=True)
    spec = load("examples/notas.yaml")
    result = Runner(spec, store=store).run()
    assert result.status == "succeeded", result.error
    rows = result.outputs["classificar"]
    assert [r["numero"] for r in rows] == ["1001", "1002", "1003"]
    assert rows[0]["cliente"] == "Acme Ltda" and rows[2]["cliente"] is None
    assert json.loads(pathlib.Path("out/notas.json").read_text())[1]["cliente"] == "Beta SA"


def test_sections_chain_and_env(tmp_path):
    spec = parse(
        {
            "name": "x",
            "extract": [{"id": "a", "use": "csv.read", "with": {"path": "${SRC}"}}],
            "transform": [{"id": "b", "use": "select", "with": {"fields": ["n"]}}],
            "load": [{"use": "json.write", "with": {"path": "${OUT:-o.json}"}}],
        },
        env={"SRC": "in.csv"},
    )
    assert spec.step("a").with_["path"] == "in.csv"
    assert spec.step("b").inputs == ["a"]
    assert spec.step("load_1").inputs == ["b"] and spec.step("load_1").with_["path"] == "o.json"


def test_missing_env_var_is_an_error():
    with pytest.raises(SpecError, match="TOKEN"):
        parse({"name": "x", "steps": [{"id": "a", "use": "http.get", "with": {"url": "${TOKEN}"}}]}, env={})


def test_spec_validation():
    with pytest.raises(SpecError, match="exactly one"):
        parse({"name": "x", "steps": [{"id": "a", "use": "csv.read", "fn": "m:f"}]})
    with pytest.raises(SpecError, match="unknown input"):
        parse({"name": "x", "steps": [{"id": "a", "use": "concat", "input": "zz"}]})
    with pytest.raises(SpecError, match="output"):
        parse({"name": "x", "steps": [{"id": "a", "agent": "do it"}]})


def test_cycle_detected():
    spec = parse(
        {
            "name": "x",
            "steps": [{"id": "a", "use": "concat", "input": "b"}, {"id": "b", "use": "concat", "input": "a"}],
        }
    )
    with pytest.raises(SpecError, match="cycle"):
        graph.order(spec)


def test_mermaid_has_nodes_edges_and_status():
    spec = parse(
        {
            "name": "x",
            "steps": [
                {"id": "a", "use": "csv.read"},
                {"id": "b", "agent": "go", "input": "a", "output": {"v": "str"}},
            ],
        }
    )
    text = graph.mermaid(spec, {"b": "failed"})
    assert "a --> b" in text and 'b{{"' in text and "class b failed" in text


def test_custom_connector_and_fn(store):
    @connector("test.numbers")
    def numbers(input, n=3):
        return [{"n": i} for i in range(n)]

    spec = parse(
        {
            "name": "x",
            "steps": [
                {"id": "src", "use": "test.numbers", "with": {"n": 4}},
                {"id": "dbl", "fn": "tests.helpers:double", "input": "src"},
            ],
        }
    )
    result = Runner(spec, store=store).run()
    assert result.status == "succeeded", result.error
    assert [r["n"] for r in result.outputs["dbl"]] == [0, 2, 4, 6]


class FlakyProvider:
    """Fails on records with bad=True, counts calls."""

    calls = 0

    def __init__(self, spec):
        pass

    def structured(self, instruction, record, schema):
        FlakyProvider.calls += 1
        if record.get("bad"):
            raise ValueError("invalid output")
        return Result(schema.model_validate({"label": f"L{record['n']}"}), 10, 2)

    def text(self, instruction, content):
        return "summary"


def test_agent_rejects_cache_and_notify(store, monkeypatch):
    register_provider("flaky", FlakyProvider)
    sent = []
    monkeypatch.setattr("agentl.notify._post", lambda url, payload: sent.append(payload))
    raw = {
        "name": "agent-test",
        "llm": {"provider": "flaky"},
        "steps": [
            {"id": "src", "use": "test.rows"},
            {"id": "label", "agent": "label it", "input": "src", "output": {"label": "str"}},
        ],
        "notify": {
            "channels": {"slack": {"webhook": "https://hooks.example/x"}},
            "rules": [{"on": "item.rejected", "when": "rate > 0.2", "to": ["slack"]}],
        },
    }

    @connector("test.rows")
    def rows(input):
        return [{"n": 1}, {"n": 2}, {"n": 3, "bad": True}]

    FlakyProvider.calls = 0
    r1 = Runner(parse(raw), store=store).run()
    assert r1.status == "succeeded"
    assert [x["label"] for x in r1.outputs["label"]] == ["L1", "L2"]
    assert len(r1.rejected["label"]) == 1 and r1.input_tokens == 20
    assert FlakyProvider.calls == 4  # 2 ok + 1 bad with one retry
    assert len(sent) == 1 and "item.rejected" in sent[0]["text"]

    r2 = Runner(parse(raw), store=store).run()
    assert FlakyProvider.calls == 6  # cached rows skipped, only the bad one retried
    assert r2.outputs["label"] == r1.outputs["label"]


def test_failed_step_marks_run_failed(store):
    spec = parse({"name": "x", "steps": [{"id": "a", "use": "csv.read", "with": {"path": "nope.csv"}}]})
    result = Runner(spec, store=store).run()
    assert result.status == "failed" and "nope.csv" in result.error
    types = [e["type"] for e in store.events(result.run_id)]
    assert types == ["run.started", "step.started", "step.failed", "run.failed"]


def test_schema_model_types():
    M = schema_model("M", {"a": "int", "b": "str?", "c": "list[str]"})
    m = M.model_validate({"a": "3", "c": ["x"]})
    assert m.a == 3 and m.b is None and m.c == ["x"]
    with pytest.raises(ValueError):
        schema_model("N", {"a": "decimal"})


def test_file_connectors(tmp_path):
    import docx
    import openpyxl

    from agentl.registry import get

    wb = openpyxl.Workbook()
    wb.active.append(["nome", "idade"])
    wb.active.append(["Ana", 30])
    wb.save(tmp_path / "p.xlsx")
    assert get("excel.read")(None, path=str(tmp_path / "p.xlsx")) == [{"nome": "Ana", "idade": 30}]

    d = docx.Document()
    d.add_paragraph("Contrato de teste")
    d.save(tmp_path / "c.docx")
    assert get("docx.read")(None, path=str(tmp_path / "*.docx"))[0]["text"] == "Contrato de teste"

    get("sqlite.write")([{"a": 1}, {"a": 2}], db=str(tmp_path / "x.db"), table="t")
    assert get("sqlite.read")(None, db=str(tmp_path / "x.db"), query="select sum(a) s from t") == [{"s": 3}]


def test_http_get_per_record(monkeypatch):
    seen = []

    class Resp:
        def __init__(self, body):
            self.body = body

        def read(self):
            return json.dumps(self.body).encode()

        def __enter__(self):
            return self

        def __exit__(self, *a):
            pass

    def fake_urlopen(req, timeout):
        seen.append(req.full_url)
        return Resp({"data": {"items": [{"id": req.full_url[-1]}]}})

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)
    from agentl.registry import get

    out = get("http.get")([{"id": 1}, {"id": 2}], url="https://api.x/c/{id}", items="data.items")
    assert sorted(seen) == ["https://api.x/c/1", "https://api.x/c/2"]
    assert sorted(r["id"] for r in out) == ["1", "2"]


def test_transforms_accept_a_single_field_name():
    from agentl.registry import get

    rows = [{"id": 1, "v": "a"}, {"id": 1, "v": "b"}, {"id": 2, "v": ""}]
    assert get("dedupe")(rows, on="id") == [{"id": 1, "v": "a"}, {"id": 2, "v": ""}]
    assert get("filter")(rows, not_empty="v") == rows[:2]
    assert get("select")(rows[:1], fields="v") == [{"v": "a"}]


def test_agent_step_with_empty_upstream_does_nothing(store):
    spec = parse(
        {
            "name": "x",
            "llm": {"provider": "echo"},
            "steps": [
                {"id": "src", "use": "filter", "input": "seed", "with": {"where": {"k": "none"}}},
                {"id": "seed", "use": "concat"},
                {"id": "ask", "agent": "Check it.", "input": "src", "output": {"k": "str?"}},
            ],
        }
    )
    result = Runner(spec, store=store).run()
    assert result.status == "succeeded", result.error
    assert result.outputs["ask"] == []


def test_invalid_yaml_is_a_spec_error(tmp_path):
    p = tmp_path / "bad.yaml"
    p.write_text("name: x\nsteps:\n  - {id: a, use: concat, output: { v: str? }}\n")
    with pytest.raises(SpecError, match="block style"):
        load(p)
