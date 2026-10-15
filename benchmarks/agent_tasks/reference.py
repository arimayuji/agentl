"""Reference solutions in plain Python: proof that every task is solvable and that the scorer accepts
a right answer. The tests use it as a stand-in agent: it reads the request on stdin, like an agent."""

import csv
import json
import re
import sqlite3
import sys
import urllib.request
from pathlib import Path

MARKERS = {
    "01": "orders.csv. Keep",
    "02": "customers.json has",
    "03": "customers.xlsx",
    "04": "/products",
    "05": "invoices/ has",
    "06": "pricing.html",
    "07": "shop.db",
    "08": "orders.json has",
    "09": "sales/ has",
    "10": "emails/ has",
}


def read_csv(path: str | Path) -> list[dict]:
    return list(csv.DictReader(Path(path).read_text().splitlines()))


def read_json(path: str) -> list[dict]:
    return json.loads(Path(path).read_text())


def write_csv(rows: list[dict], cols: list[str]) -> None:
    with open("out/result.csv", "w", newline="") as f:
        w = csv.DictWriter(f, cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_json(rows: list[dict]) -> None:
    Path("out/result.json").write_text(json.dumps(rows))


def main() -> None:
    prompt = sys.stdin.read()
    task = next(k for k, m in MARKERS.items() if m in prompt)
    Path("out").mkdir(exist_ok=True)
    if task == "01":
        rows = [r for r in read_csv("input/orders.csv") if r["status"] == "paid" and float(r["total"]) > 100]
        write_json(
            [{"id": int(r["id"]), "customer": r["customer"], "total": float(r["total"])} for r in rows]
        )
    elif task == "02":
        names = {c["customer_id"]: c["name"] for c in read_json("input/customers.json")}
        rows = [
            {
                "order_id": r["order_id"],
                "customer_name": names.get(int(r["customer_id"]), ""),
                "total": r["total"],
            }
            for r in read_csv("input/orders.csv")
        ]
        write_csv(rows, ["order_id", "customer_name", "total"])
    elif task == "03":
        from openpyxl import load_workbook

        seen: dict = {}
        sheet = load_workbook("input/customers.xlsx")["export"]
        for row in list(sheet.iter_rows(values_only=True))[1:]:
            seen.setdefault(row[0], {"id": row[0], "name": row[1], "email": row[2]})
        write_csv(list(seen.values()), ["id", "name", "email"])
    elif task == "04":
        url = re.search(r"(http://127\.0\.0\.1:\d+)", prompt).group(1)  # type: ignore[union-attr]
        page, out = 1, []
        while page:
            data = json.load(urllib.request.urlopen(f"{url}/products?page={page}"))
            out += data["data"]
            page = data["next_page"]
        write_json(out)
    elif task == "05":
        from pypdf import PdfReader

        rows = []
        for p in sorted(Path("input/invoices").glob("*.pdf")):
            text = PdfReader(p).pages[0].extract_text()
            number = re.search(r"INVOICE (\d+)", text).group(1)  # type: ignore[union-attr]
            total = re.search(r"R\$ ([\d.,]+)", text).group(1)  # type: ignore[union-attr]
            rows.append({"number": number, "total": total.replace(".", "").replace(",", ".")})
        write_csv(rows, ["number", "total"])
    elif task == "06":
        html = Path("input/pricing.html").read_text()
        cells = re.findall(r"<tr><td>(.*?)</td><td>(.*?)</td><td>(.*?)</td></tr>", html)
        write_csv([{"plan": a, "seats": b, "price": c} for a, b, c in cells], ["plan", "seats", "price"])
    elif task == "07":
        con = sqlite3.connect("input/shop.db")
        query = "select region, count(*), sum(total) from orders group by region"
        write_json([{"region": r, "orders": n, "total": t} for r, n, t in con.execute(query)])
    elif task == "08":
        rows = [
            {
                "id": o["id"],
                "customer_name": o["customer"]["name"],
                "city": o["customer"]["city"],
                "item_count": len(o["items"]),
                "quantity": sum(i["qty"] for i in o["items"]),
            }
            for o in read_json("input/orders.json")
        ]
        write_csv(rows, ["id", "customer_name", "city", "item_count", "quantity"])
    elif task == "09":
        rows = [r for p in sorted(Path("input/sales").glob("*.csv")) for r in read_csv(p)]
        write_csv(rows, ["date", "store", "total"])
    elif task == "10":
        out = []
        for p in sorted(Path("input/emails").glob("*.txt")):
            m = re.search(r"order #?(\d+)", p.read_text())
            out.append({"file": p.name, "order": m.group(1) if m else None})
        write_json(out)
    # what Claude Code prints with --output-format json, so the metrics path is exercised too
    print(
        json.dumps({"total_cost_usd": 0.0, "num_turns": 1, "usage": {"input_tokens": 10, "output_tokens": 5}})
    )


if __name__ == "__main__":
    main()
