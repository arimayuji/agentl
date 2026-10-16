"""File connectors.

Every reader takes its files from `path` (a glob), `url` (downloaded first) or the
`path` field of upstream records, and is built on a per-file function so that
`file.read` can dispatch mixed folders and archives to the right one.
"""

from __future__ import annotations

import csv
import datetime as _dt
import gzip
import json
import logging
import re
import shutil
import subprocess
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Callable, Iterable, Iterator
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from ..registry import arrow_reader, arrow_writer, connector, require, stream_reader, stream_writer
from ._util import is_remote, origin, paths, single, tmpdir, upload

log = logging.getLogger("agentl.files")

Records = list[dict]


# ---------------------------------------------------------------- per-file readers
def iter_csv(p: Path, delimiter: str | None = None, encoding: str = "utf-8", **_: Any) -> Iterator[dict]:
    with p.open(newline="", encoding=encoding, errors="replace") as f:
        sample = f.read(64 * 1024)
        f.seek(0)
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = "\t" if p.suffix.lower() == ".tsv" else ","
        # csv.reader and zip: the same records as csv.DictReader (short rows padded with None,
        # extra cells under the key None), about twice as fast.
        reader = csv.reader(f, delimiter=delimiter)
        header = next(reader, None)
        if header is None:
            return
        width = len(header)
        for cells in reader:
            if len(cells) == width:
                yield dict(zip(header, cells, strict=True))
            elif not cells:
                continue  # DictReader skips blank lines too
            elif len(cells) < width:
                yield dict(zip(header, cells + [None] * (width - len(cells)), strict=True))
            else:
                row: dict[Any, Any] = dict(zip(header, cells[:width], strict=True))
                row[None] = cells[width:]
                yield row


def read_csv(p: Path, delimiter: str | None = None, encoding: str = "utf-8", **_: Any) -> Records:
    return list(iter_csv(p, delimiter, encoding))


def read_json(p: Path, key: str | None = None, encoding: str = "utf-8", **_: Any) -> Records:
    text = p.read_text(encoding=encoding)
    if p.suffix.lower() in (".jsonl", ".ndjson"):
        return [json.loads(line) for line in text.splitlines() if line.strip()]
    data = json.loads(text)
    if key:
        for part in key.split("."):
            data = data[part]
    rows = data if isinstance(data, list) else [data]
    return [r if isinstance(r, dict) else {"value": r} for r in rows]


def read_txt(p: Path, lines: bool = False, encoding: str = "utf-8", **_: Any) -> Records:
    text = p.read_text(encoding=encoding, errors="replace")
    if lines:
        return [
            {"path": str(p), "line": i + 1, "text": t} for i, t in enumerate(text.splitlines()) if t.strip()
        ]
    return [{"path": str(p), "text": text}]


def read_xlsx(p: Path, sheet: str | None = None, header_row: int = 1, **_: Any) -> Records:
    openpyxl = require("openpyxl", "excel")
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.active
        rows = ws.iter_rows(min_row=header_row, values_only=True)
        header = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(next(rows, []))]
        return [dict(zip(header, row, strict=False)) for row in rows if any(v is not None for v in row)]
    finally:
        wb.close()


def read_xls(p: Path, sheet: str | None = None, header_row: int = 1, **_: Any) -> Records:
    xlrd = require("xlrd", "xls")
    book = xlrd.open_workbook(str(p))
    ws = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
    start = header_row - 1
    header = [str(v) if v != "" else f"col_{i}" for i, v in enumerate(ws.row_values(start))]
    out = []
    for r in range(start + 1, ws.nrows):
        row = ws.row_values(r)
        if any(v != "" for v in row):
            out.append(dict(zip(header, row, strict=False)))
    return out


def read_parquet(p: Path, columns: list[str] | None = None, **_: Any) -> Records:
    pq = require("pyarrow.parquet", "parquet")
    return list(pq.read_table(str(p), columns=columns).to_pylist())


def read_pdf(p: Path, **_: Any) -> Records:
    pypdf = require("pypdf", "pdf")
    reader = pypdf.PdfReader(str(p))
    pages = [page.extract_text() or "" for page in reader.pages]
    return [{"path": str(p), "pages": len(pages), "text": "\n".join(pages)}]


def read_docx(p: Path, **_: Any) -> Records:
    docx = require("docx", "office")
    doc = docx.Document(str(p))
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
    return [{"path": str(p), "text": "\n".join(paragraphs), "tables": tables}]


def read_doc(p: Path, **opts: Any) -> Records:
    """Legacy Word files: convert with LibreOffice, then read as .docx."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError(f"{p.name}: reading .doc needs LibreOffice ('soffice') on the PATH")
    out_dir = tmpdir("doc-")
    subprocess.run(
        [soffice, "--headless", "--convert-to", "docx", "--outdir", str(out_dir), str(p)],
        check=True,
        capture_output=True,
        timeout=120,
    )
    records = read_docx(out_dir / (p.stem + ".docx"), **opts)
    for r in records:
        r["path"] = str(p)
    return records


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]  # drop XML namespace


def read_xml(p: Path, items: str | None = None, **_: Any) -> Records:
    """One record per `items` element (default: children of the root).

    Attributes and simple child elements become fields.
    """
    data = p.read_bytes()
    # The stdlib parser never fetches external entities, but it does expand internal ones, so a
    # small file can blow up into gigabytes ("billion laughs"). Data files do not need them.
    if b"<!ENTITY" in data:
        raise ValueError(f"{p.name}: XML with entity declarations (<!ENTITY) is not read")
    root = ET.fromstring(data)
    nodes = [e for e in root.iter() if _local(e.tag) == items] if items else list(root)
    out = []
    for node in nodes:
        rec: dict[str, Any] = {_local(k): v for k, v in node.attrib.items()}
        for child in node:
            if len(child) == 0:
                rec[_local(child.tag)] = (child.text or "").strip()
        if not rec and (node.text or "").strip():
            rec[_local(node.tag)] = node.text.strip()  # type: ignore[union-attr]
        out.append(rec)
    return out


class _Tables(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.tables: list[list[list[str]]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: Any) -> None:
        if tag == "table":
            self.tables.append([])
        elif tag == "tr" and self.tables:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if self._row:
                self.tables[-1].append(self._row)
            self._row = None

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)


def read_html(p: Path, table: int | None = None, encoding: str = "utf-8", **_: Any) -> Records:
    """Rows of the HTML tables; the first row of each table is the header."""
    return html_tables(p.read_text(encoding=encoding, errors="replace"), table)


def html_tables(html: str, table: int | None = None) -> Records:
    """Rows of the tables in an HTML string; the first row of each table is the header."""
    parser = _Tables()
    parser.feed(html)
    out: Records = []
    for i, rows in enumerate(parser.tables):
        if (table is not None and i != table) or len(rows) < 2:
            continue
        header = [h or f"col_{j}" for j, h in enumerate(rows[0])]
        out.extend({"table": i, **dict(zip(header, r, strict=False))} for r in rows[1:])
    return out


READERS: dict[str, Callable[..., Records]] = {
    ".csv": read_csv,
    ".tsv": read_csv,
    ".json": read_json,
    ".jsonl": read_json,
    ".ndjson": read_json,
    ".txt": read_txt,
    ".md": read_txt,
    ".log": read_txt,
    ".xlsx": read_xlsx,
    ".xlsm": read_xlsx,
    ".xls": read_xls,
    ".parquet": read_parquet,
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".doc": read_doc,
    ".xml": read_xml,
    ".html": read_html,
    ".htm": read_html,
}


# ---------------------------------------------------------------- archives
MAX_EXTRACT_BYTES = 1024**3  # 1 GiB per archive, guards against zip bombs


def _extract_zip(p: Path, limit: int) -> list[Path]:
    dest = tmpdir("zip-")
    out = []
    written = 0
    with zipfile.ZipFile(p) as z:
        infos = [i for i in z.infolist() if not i.is_dir()]
        if sum(i.file_size for i in infos) > limit:
            raise ValueError(f"{p.name}: archive expands beyond {limit} bytes")
        for info in infos:
            target = (dest / info.filename).resolve()
            if not target.is_relative_to(dest.resolve()):
                raise ValueError(f"{p.name}: unsafe path in archive: {info.filename}")
            target.parent.mkdir(parents=True, exist_ok=True)
            with z.open(info) as src, target.open("wb") as dst:
                # count what is really written: the sizes in a zip header can lie
                while chunk := src.read(1024 * 1024):
                    written += len(chunk)
                    if written > limit:
                        raise ValueError(f"{p.name}: archive expands beyond {limit} bytes")
                    dst.write(chunk)
            out.append(target)
    return sorted(out)


def _extract_gz(p: Path, limit: int) -> list[Path]:
    target = tmpdir("gz-") / p.stem  # data.csv.gz -> data.csv
    written = 0
    with gzip.open(p, "rb") as src, target.open("wb") as dst:
        while chunk := src.read(1024 * 1024):
            written += len(chunk)
            if written > limit:
                raise ValueError(f"{p.name}: archive expands beyond {limit} bytes")
            dst.write(chunk)
    return [target]


def _expand(p: Path, limit: int, depth: int = 0) -> list[tuple[Path, Path | None]]:
    """(file, archive it came from) pairs, expanding zip/gz up to two levels deep."""
    suffix = p.suffix.lower()
    if suffix in (".zip", ".gz") and depth < 2:
        inner = _extract_zip(p, limit) if suffix == ".zip" else _extract_gz(p, limit)
        return [(f, p) for x in inner for f, _ in _expand(x, limit, depth + 1)]
    return [(p, None)]


# ---------------------------------------------------------------- connectors
def _each(
    reader: Callable[..., Records],
    input: Any,
    path: str | None,
    url: str | None,
    headers: dict | None,
    field: str = "path",
    **opts: Any,
) -> Records:
    out: Records = []
    so = opts.pop("storage_options", None)
    for p in paths(input, path, field, url, headers, so):
        out.extend(reader(p, **opts))
    return out


@connector("file.read")
def file_read(
    input: Any = None,
    path: str | None = None,
    url: str | None = None,
    headers: dict | None = None,
    field: str = "path",
    on_unsupported: str = "skip",
    max_extract_bytes: int = MAX_EXTRACT_BYTES,
    storage_options: dict | None = None,
    **opts: Any,
) -> Records:
    """Read any supported file, detecting the format by extension.

    ZIP and GZ archives are expanded and each inner file is read with its own reader.
    Each record gets `_file` (and `_archive` when it came from one).
    """
    if on_unsupported not in ("skip", "error"):
        raise ValueError("on_unsupported must be 'skip' or 'error'")
    out: Records = []
    for p in paths(input, path, field, url, headers, storage_options):
        for f, archive in _expand(p, max_extract_bytes):
            reader = READERS.get(f.suffix.lower())
            if reader is None:
                if on_unsupported == "error":
                    raise ValueError(f"unsupported file type: {f.name}")
                log.warning("file.read: skipping unsupported file %s", f.name)
                continue
            shown = f"{archive.name}!{f.name}" if archive else origin(f)
            for rec in reader(f, **opts):
                rec.setdefault("_file", shown)
                if archive:
                    rec.setdefault("_archive", origin(archive))
                out.append(rec)
    return out


@connector("csv.read")
def csv_read(
    input: Any = None,
    path: str | None = None,
    url: str | None = None,
    headers: dict | None = None,
    storage_options: dict | None = None,
    delimiter: str | None = None,
    encoding: str = "utf-8",
) -> Records:
    return _each(
        read_csv,
        input,
        path,
        url,
        headers,
        delimiter=delimiter,
        encoding=encoding,
        storage_options=storage_options,
    )


@stream_reader("csv.read")
def csv_read_chunks(
    input: Any = None,
    chunk_size: int = 50_000,
    path: str | None = None,
    url: str | None = None,
    headers: dict | None = None,
    storage_options: dict | None = None,
    delimiter: str | None = None,
    encoding: str = "utf-8",
) -> Iterator[Records]:
    chunk: Records = []
    for p in paths(input, path, "path", url, headers, storage_options):
        for row in iter_csv(p, delimiter, encoding):
            chunk.append(row)
            if len(chunk) >= chunk_size:
                yield chunk
                chunk = []
    if chunk:
        yield chunk


def _sniff(p: Path, delimiter: str | None, encoding: str) -> tuple[str, list[str], bool]:
    """(delimiter, header, starts with a BOM) the way `iter_csv` sees them."""
    with p.open(newline="", encoding=encoding, errors="replace") as f:
        sample = f.read(64 * 1024)
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = "\t" if p.suffix.lower() == ".tsv" else ","
        header = next(csv.reader(sample.splitlines() or [""], delimiter=delimiter), [])
    return delimiter, header, sample.startswith("\ufeff")


@arrow_reader("csv.read")
def csv_read_arrow(
    input: Any = None,
    path: str | None = None,
    url: str | None = None,
    headers: dict | None = None,
    storage_options: dict | None = None,
    delimiter: str | None = None,
    encoding: str = "utf-8",
    _engine: str | None = None,
) -> Any:
    """pyarrow's CSV reader with every column as text, like `csv.read`. Falls back to the normal
    reader (and returns records) for anything it would read differently: several files, another
    encoding, a BOM, duplicate or ragged columns, invalid UTF-8. With `_engine="polars"` (the runner
    sets it when every consumer is a polars step) polars reads the file, falling back to pyarrow."""
    files = paths(input, path, "path", url, headers, storage_options)
    if len(files) != 1 or encoding.lower().replace("-", "") != "utf8":
        return csv_read(input, path, url, headers, storage_options, delimiter, encoding)
    pa = require("pyarrow", "arrow")
    pacsv = require("pyarrow.csv", "arrow")
    p = files[0]
    sep, header, bom = _sniff(p, delimiter, encoding)
    if bom or not header or len(set(header)) != len(header) or "\n" in "".join(header):
        return [r for r in iter_csv(p, delimiter, encoding)]
    if _engine == "polars":
        table = _csv_read_polars(p, sep, header)
        if table is not None:
            return table
    try:
        table = pacsv.read_csv(
            str(p),
            read_options=pacsv.ReadOptions(use_threads=True),
            parse_options=pacsv.ParseOptions(delimiter=sep, newlines_in_values=True),
            convert_options=pacsv.ConvertOptions(
                column_types={h: pa.string() for h in header},
                strings_can_be_null=False,
                quoted_strings_can_be_null=False,
                null_values=[],
            ),
        )
    except pa.ArrowInvalid:
        return [r for r in iter_csv(p, delimiter, encoding)]
    if table.column_names != header or any(t != pa.string() for t in table.schema.types):
        return [r for r in iter_csv(p, delimiter, encoding)]
    return table


def _csv_read_polars(p: Path, sep: str, header: list[str]) -> Any:
    """polars' CSV reader with every column as text and empty fields as "" (as `csv.read`), or
    None when polars reads the file differently or fails, so the caller falls back."""
    try:
        pl = require("polars", "polars")
        df = pl.read_csv(str(p), separator=sep, infer_schema=False, truncate_ragged_lines=False)
    except Exception:
        return None
    if df.columns != header:
        return None
    # Only polars steps read this table, so polars' own string type stays (no cast both ways).
    return df.with_columns(pl.all().fill_null("")).to_arrow()


@connector("json.read")
def json_read(
    input: Any = None,
    path: str | None = None,
    url: str | None = None,
    headers: dict | None = None,
    storage_options: dict | None = None,
    key: str | None = None,
) -> Records:
    return _each(read_json, input, path, url, headers, key=key, storage_options=storage_options)


@connector("txt.read")
def txt_read(
    input: Any = None,
    path: str | None = None,
    url: str | None = None,
    headers: dict | None = None,
    storage_options: dict | None = None,
    lines: bool = False,
    encoding: str = "utf-8",
) -> Records:
    return _each(
        read_txt, input, path, url, headers, lines=lines, encoding=encoding, storage_options=storage_options
    )


@connector("excel.read")
def excel_read(
    input: Any = None,
    path: str | None = None,
    url: str | None = None,
    headers: dict | None = None,
    storage_options: dict | None = None,
    sheet: str | None = None,
    header_row: int = 1,
) -> Records:
    def read(p: Path, **kw: Any) -> Records:
        return (read_xls if p.suffix.lower() == ".xls" else read_xlsx)(p, **kw)

    return _each(
        read, input, path, url, headers, sheet=sheet, header_row=header_row, storage_options=storage_options
    )


@connector("parquet.read")
def parquet_read(
    input: Any = None,
    path: str | None = None,
    url: str | None = None,
    headers: dict | None = None,
    storage_options: dict | None = None,
    columns: list[str] | None = None,
) -> Records:
    return _each(read_parquet, input, path, url, headers, columns=columns, storage_options=storage_options)


@arrow_reader("parquet.read")
def parquet_read_arrow(
    input: Any = None,
    path: str | None = None,
    url: str | None = None,
    headers: dict | None = None,
    storage_options: dict | None = None,
    columns: list[str] | None = None,
) -> Any:
    pa = require("pyarrow", "arrow")
    pq = require("pyarrow.parquet", "arrow")
    tables = [
        pq.read_table(str(p), columns=columns)
        for p in paths(input, path, "path", url, headers, storage_options)
    ]
    if not tables:
        return []
    if len(tables) > 1 and any(t.schema != tables[0].schema for t in tables):
        return [r for t in tables for r in t.to_pylist()]  # same as the normal reader
    return pa.concat_tables(tables)


@connector("pdf.read")
def pdf_read(
    input: Any = None,
    path: str | None = None,
    url: str | None = None,
    headers: dict | None = None,
    storage_options: dict | None = None,
    field: str = "path",
) -> Records:
    return _each(read_pdf, input, path, url, headers, field, storage_options=storage_options)


@connector("docx.read")
def docx_read(
    input: Any = None,
    path: str | None = None,
    url: str | None = None,
    headers: dict | None = None,
    storage_options: dict | None = None,
    field: str = "path",
) -> Records:
    return _each(read_docx, input, path, url, headers, field, storage_options=storage_options)


@connector("doc.read")
def doc_read(
    input: Any = None,
    path: str | None = None,
    url: str | None = None,
    headers: dict | None = None,
    storage_options: dict | None = None,
    field: str = "path",
) -> Records:
    return _each(read_doc, input, path, url, headers, field, storage_options=storage_options)


@connector("xml.read")
def xml_read(
    input: Any = None,
    path: str | None = None,
    url: str | None = None,
    headers: dict | None = None,
    storage_options: dict | None = None,
    items: str | None = None,
) -> Records:
    return _each(read_xml, input, path, url, headers, items=items, storage_options=storage_options)


@connector("html.read")
def html_read(
    input: Any = None,
    path: str | None = None,
    url: str | None = None,
    headers: dict | None = None,
    storage_options: dict | None = None,
    table: int | None = None,
    format: str = "tables",
    evidence: bool | str = False,
    evidence_keep_days: int | None = None,
    retry_if_empty: int = 0,
    encoding: str = "utf-8",
) -> Records:
    """`format: tables` (default) gives the rows of the HTML tables. `format: markdown` gives one
    record per page with `title` and `markdown`: the main content, without menus, footers or
    scripts (needs `agentl-etl[web]`). `evidence` keeps the raw HTML and the Markdown of every page
    (true: `.agentl/evidence`, or a folder); `retry_if_empty` reads a page again when it gave
    nothing."""
    if format not in ("tables", "markdown"):
        raise ValueError("format must be 'tables' or 'markdown'")
    if format == "tables" and not evidence and not retry_if_empty:
        return _each(read_html, input, path, url, headers, table=table, storage_options=storage_options)
    from . import web

    if evidence and evidence_keep_days:
        web.prune_evidence(evidence, evidence_keep_days)
    out: Records = []
    for attempt in range(retry_if_empty + 1):
        out = []
        empty = False
        for p in paths(input, path, "path", url, headers, storage_options):
            html = p.read_text(encoding=encoding, errors="replace")
            source = url if url and "{" not in url else str(p)
            if format == "markdown":
                page = web.to_markdown(html, url=source if source.startswith("http") else None)
                rows = [{**page, "_source": source}] if page["markdown"] else []
            else:
                rows = html_tables(html, table)
            empty = empty or not rows
            ev = (
                web.save_evidence(evidence, source, html, page["markdown"] if format == "markdown" else None)
                if evidence
                else {}
            )
            out.extend({**r, **ev} for r in rows)
            if not rows and format == "markdown" and attempt == retry_if_empty:
                out.append({"title": "", "markdown": "", "_source": source, **ev})
        if not empty:
            break
    return out


# ---------------------------------------------------------------- writers
def _target(path: str) -> Path:
    """Local file to write: the path itself, or a temp file uploaded afterwards."""
    local = tmpdir("out-") / Path(path).name if is_remote(path) else Path(path)
    local.parent.mkdir(parents=True, exist_ok=True)
    return local


# Versioned outputs: `{date}` / `{datetime}` in the path, a `latest` copy, `keep` the newest N.
_STAMPS = {"date": "%Y-%m-%d", "datetime": "%Y-%m-%dT%H%M%S"}


def versioned(path: str, now: _dt.datetime | None = None) -> str:
    """'reports/sales_{date}.xlsx' -> 'reports/sales_2026-10-14.xlsx' (local time)."""
    if "{" not in path:
        return path
    now = now or _dt.datetime.now()
    try:
        return path.format(**{k: now.strftime(f) for k, f in _STAMPS.items()})
    except (KeyError, IndexError) as e:
        raise ValueError(f"path '{path}': only {{date}} and {{datetime}} can be used") from e


def publish(
    template: str, path: str, local: Path, latest: str | None, keep: int | None, storage_options: dict | None
) -> dict:
    """After a writer wrote `local` to `path`: copy it to `latest` and delete all but the newest
    `keep` files made from the same `template`. Returns the fields to add to the writer's result."""
    extra: dict = {}
    if latest:
        if is_remote(latest):
            upload(local, latest, storage_options)
        else:
            Path(latest).parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(local, latest)
        extra["latest"] = latest
    if keep:
        if "{" not in template:
            raise ValueError("keep needs {date} or {datetime} in the path, so versions have their own names")
        pattern = template.format(**dict.fromkeys(_STAMPS, "*"))
        # Only names in the template's own format count as versions (not sales_latest.csv).
        name = re.escape(template.replace("\\", "/").rsplit("/", 1)[-1])
        for key, regex in (("date", r"\d{4}-\d{2}-\d{2}"), ("datetime", r"\d{4}-\d{2}-\d{2}T\d{6}")):
            name = name.replace(re.escape("{" + key + "}"), regex)
        version = re.compile(name + "$")

        def versions(paths: list[str]) -> list[str]:
            return sorted(p for p in paths if version.match(p.replace("\\", "/").rsplit("/", 1)[-1]))

        if is_remote(template):
            from ._util import _fs

            fs, root = _fs(pattern, storage_options)
            found = versions(fs.glob(root))
            for old in found[:-keep]:
                fs.rm(old)
        else:
            import glob as _glob

            found = versions(_glob.glob(pattern))
            for old in found[:-keep]:
                Path(old).unlink(missing_ok=True)
        extra["deleted"] = max(0, len(found) - keep)
    return extra


@stream_writer("csv.write")
def csv_write_chunks(
    chunks: Iterable[Records],
    path: str,
    delimiter: str = ",",
    storage_options: dict | None = None,
    latest: str | None = None,
    keep: int | None = None,
) -> Records:
    """The header comes from the first chunk with records. A column that first appears in a later
    chunk is an error rather than silently dropped: select the fields, or use a bigger chunk."""
    template, path = path, versioned(path)
    local = _target(path)
    written = 0
    w: csv.DictWriter | None = None
    with open(local, "w", newline="", encoding="utf-8") as f:
        for rows in chunks:
            if w is None:
                if not rows:
                    continue
                fields: list[str] = []
                for r in rows:
                    fields.extend(k for k in r if k not in fields)
                w = csv.DictWriter(f, fieldnames=fields, delimiter=delimiter)
                w.writeheader()
            try:
                w.writerows(rows)
            except ValueError as e:
                raise ValueError(f"csv.write: {e} (column not in the first chunk)") from e
            written += len(rows)
        if w is None:
            csv.writer(f).writerow([])
    if is_remote(path):
        upload(local, path, storage_options)
    return [
        {"path": path, "written": written, **publish(template, path, local, latest, keep, storage_options)}
    ]


@connector("csv.write")
def csv_write(
    input: Any,
    path: str,
    delimiter: str = ",",
    storage_options: dict | None = None,
    latest: str | None = None,
    keep: int | None = None,
) -> Records:
    return csv_write_chunks([single(input)], path, delimiter, storage_options, latest, keep)


def _write_file(
    path: str,
    storage_options: dict | None,
    write: Callable[[Path], None],
    latest: str | None,
    keep: int | None,
    written: int,
) -> Records:
    """Writers that build the whole file at once: write locally, upload, publish versions."""
    template, path = path, versioned(path)
    local = _target(path)
    write(local)
    if is_remote(path):
        upload(local, path, storage_options)
    return [
        {"path": path, "written": written, **publish(template, path, local, latest, keep, storage_options)}
    ]


@connector("parquet.write")
def parquet_write(
    input: Any,
    path: str,
    compression: str = "snappy",
    storage_options: dict | None = None,
    latest: str | None = None,
    keep: int | None = None,
) -> Records:
    """Write records to a Parquet file (local or remote). Needs `agentl-etl[parquet]`."""
    pa = require("pyarrow", "parquet")
    pq = require("pyarrow.parquet", "parquet")
    rows = single(input)
    return _write_file(
        path,
        storage_options,
        lambda p: pq.write_table(pa.Table.from_pylist(rows), str(p), compression=compression),
        latest,
        keep,
        len(rows),
    )


@arrow_writer("parquet.write")
def parquet_write_arrow(
    data: Any,
    path: str,
    compression: str = "snappy",
    storage_options: dict | None = None,
    latest: str | None = None,
    keep: int | None = None,
) -> Records:
    from ..polars import from_polars, write_parquet

    if from_polars(data):
        write = lambda p: write_parquet(data, str(p), compression)  # noqa: E731
    else:
        pq = require("pyarrow.parquet", "arrow")
        write = lambda p: pq.write_table(data, str(p), compression=compression)  # noqa: E731
    return _write_file(
        path,
        storage_options,
        write,
        latest,
        keep,
        data.num_rows,
    )


@stream_writer("json.write")
def json_write_chunks(
    chunks: Iterable[Records],
    path: str,
    storage_options: dict | None = None,
    latest: str | None = None,
    keep: int | None = None,
) -> Records:
    """`.jsonl` is written chunk by chunk; a `.json` array needs every record, so it is collected."""
    if not path.endswith(".jsonl"):
        return json_write([r for rows in chunks for r in rows], path, storage_options, latest, keep)
    template, path = path, versioned(path)
    local = _target(path)
    written = 0
    with open(local, "w", encoding="utf-8") as f:
        for rows in chunks:
            f.writelines(json.dumps(r, ensure_ascii=False, default=str) + "\n" for r in rows)
            written += len(rows)
    if is_remote(path):
        upload(local, path, storage_options)
    return [
        {"path": path, "written": written, **publish(template, path, local, latest, keep, storage_options)}
    ]


@connector("json.write")
def json_write(
    input: Any,
    path: str,
    storage_options: dict | None = None,
    latest: str | None = None,
    keep: int | None = None,
) -> Records:
    rows = single(input)
    if path.endswith(".jsonl"):
        return json_write_chunks([rows], path, storage_options, latest, keep)

    def write(local: Path) -> None:
        with open(local, "w", encoding="utf-8") as f:
            json.dump(rows, f, ensure_ascii=False, indent=2, default=str)

    return _write_file(path, storage_options, write, latest, keep, len(rows))


# ---------------------------------------------------------------- excel.write
_NUMBER = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")


def _cell(v: Any, numbers: bool) -> Any:
    """A value openpyxl can store: JSON for dicts and lists, naive datetimes (Excel has no time
    zones), and numeric text as numbers unless it has a leading zero (codes like 00123 stay text)."""
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False, default=str)
    if isinstance(v, _dt.datetime) and v.tzinfo is not None:
        return v.astimezone(_dt.timezone.utc).replace(tzinfo=None)
    if numbers and isinstance(v, str) and _NUMBER.match(v.strip()):
        n = float(v)
        return int(n) if n.is_integer() and "." not in v else n
    return v


def _sheet_title(name: str, used: set[str]) -> str:
    title = re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "Sheet"
    base, n = title, 2
    while title.lower() in used:
        title = f"{base[:28]}_{n}"
        n += 1
    used.add(title.lower())
    return title


@connector("excel.write")
def excel_write(
    input: Any,
    path: str,
    sheet: str = "data",
    numbers: bool = True,
    freeze_header: bool = True,
    autofilter: bool = True,
    storage_options: dict | None = None,
    latest: str | None = None,
    keep: int | None = None,
) -> Records:
    """Write an .xlsx a person can open as is: bold header, frozen first row, a filter on the
    header, column widths that fit, numbers as numbers. With several inputs, one sheet per input
    step (`input: [summary, detail]`). Needs `agentl-etl[excel]`."""
    require("openpyxl", "excel")
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheets = input if isinstance(input, dict) else {sheet: input or []}
    total = sum(len(rows) for rows in sheets.values())

    def write(local: Path) -> None:
        wb = Workbook(write_only=True)
        used: set[str] = set()
        for name, rows in sheets.items():
            ws = wb.create_sheet(_sheet_title(str(name), used))
            cols: list[str] = []
            for r in rows:
                cols.extend(k for k in r if k not in cols)
            for i, c in enumerate(cols, start=1):
                longest = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:500]])
                ws.column_dimensions[get_column_letter(i)].width = min(60, longest + 2)
            if freeze_header:
                ws.freeze_panes = "A2"
            if autofilter and cols:
                ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
            header = []
            for c in cols:
                cell = WriteOnlyCell(ws, value=c)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E8EEF5")
                header.append(cell)
            ws.append(header)
            for r in rows:
                ws.append([_cell(r.get(c), numbers) for c in cols])
        wb.save(local)

    return _write_file(path, storage_options, write, latest, keep, total)
