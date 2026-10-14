"""Triage support tickets with an LLM and hand the team a spreadsheet."""

import csv
import json
from pathlib import Path

TITLE = "Support ticket triage"
NEEDS: tuple[str, ...] = ("pandas", "openpyxl")
OUTPUT = "out/triage.xlsx"


def make_data(work: Path) -> None:
    rows = [
        json.loads(x)
        for x in (Path(__file__).parents[2] / "evals/data/tickets.jsonl").read_text().splitlines()
    ]
    with open(work / "data/tickets.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["id", "text", "created_at"])
        for r in rows:
            w.writerow([r["id"], r["text"], "2026-10-14"])


def result(work: Path) -> list:
    import openpyxl

    ws = openpyxl.load_workbook(work / OUTPUT).active
    head = [c.value for c in ws[1]]
    rows = [dict(zip(head, (c.value for c in r), strict=False)) for r in ws.iter_rows(min_row=2)]
    return sorted((str(r["id"]), r["category"], bool(r["urgent"])) for r in rows)
